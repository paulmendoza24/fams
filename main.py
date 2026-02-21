import warnings

warnings.filterwarnings("ignore", category=UserWarning)
import os, sys, subprocess, threading
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
from docxtpl import DocxTemplate
from docxcompose.composer import Composer
from PIL import Image, ImageTk
from docx import Document
from PyPDF2 import PdfMerger
import unicodedata
import time
from PyPDF2 import PdfReader
import pygame
import pyttsx3

import customtkinter as ctk

BASE_OUT = "fams_output"
DOCX_OUT = os.path.join(BASE_OUT, "docx")
PDF_OUT = os.path.join(BASE_OUT, "pdf")
MERGED_DOCX_OUT = os.path.join(BASE_OUT, "merged_docx")
MERGED_PDF_OUT = os.path.join(BASE_OUT, "merged_pdf")
LOG_FILE = os.path.join(BASE_OUT, "fams_log.txt")
encodings = [
    "utf-8",
    "utf-8-sig",
    "latin1",
    "cp1252",
    "ascii",
    "utf-16",
    "utf-16-le",
    "utf-16-be",
    "cp850",
    "cp437",
    "mac_roman",
]
green_hex = "#228B22"
hover_green_hex = "#1F7E1F"


def wait_for_pdf(path, timeout=15):
    start = time.time()
    while time.time() - start < timeout:
        if os.path.exists(path) and os.path.getsize(path) > 1024:
            return True
        time.sleep(0.3)
    return False


def play_wav(file_path):
    pygame.mixer.init()
    pygame.mixer.music.load(file_path)
    pygame.mixer.music.play()

    while pygame.mixer.music.get_busy():
        continue


def is_valid_pdf(path):
    try:
        reader = PdfReader(path)
        return len(reader.pages) > 0
    except Exception:
        return False


def sanitize_filename(name):
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    return "".join(c if c.isalnum() or c in " -_." else "_" for c in name)


def read_students(path, log_func=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in [".csv", ".txt"]:
        for enc in encodings:
            try:
                df = pd.read_csv(path, dtype=str, encoding=enc)
                if log_func:
                    log_func(f"Successfully read file with encoding: {enc}")
                break
            except UnicodeDecodeError:
                if log_func:
                    log_func(f"Failed to read file with encoding: {enc}")
        else:
            err_msg = "Failed to decode file with all tried encodings."
            if log_func:
                log_func(err_msg)
            messagebox.showerror("Encoding Error", err_msg)
            raise ValueError(err_msg)
    elif ext in [".xls", ".xlsx"]:
        try:
            df = pd.read_excel(path, dtype=str)
            if log_func:
                log_func("Successfully read Excel file.")
        except Exception as e:
            err_msg = f"Failed to read Excel file: {e}"
            if log_func:
                log_func(err_msg)
            messagebox.showerror("File Read Error", err_msg)
            raise
    else:
        err_msg = "Unsupported file type. Use CSV or Excel."
        if log_func:
            log_func(err_msg)
        messagebox.showerror("File Type Error", err_msg)
        raise ValueError(err_msg)

    # ---------------- Column Detection ----------------
    cols = {c.strip().lower(): c for c in df.columns}
    name_col, id_col = None, None
    for cand in ["name", "full name", "student name", "student"]:
        if cand in cols:
            name_col = cols[cand]
            break
    for cand in [
        "student_number",
        "student no",
        "student_no",
        "id",
        "studentid",
        "student number",
    ]:
        if cand in cols:
            id_col = cols[cand]
            break

    # ---------------- Error if placeholders missing ----------------
    if name_col is None or id_col is None:
        err_msg = "❌ Missing required columns in student file! Expected 'name' and 'student_number'."
        if log_func:
            log_func(err_msg)
        messagebox.showerror("Column Error", err_msg)
        raise ValueError(err_msg)

    # ---------------- Load Students ----------------
    students = []
    for _, row in df.iterrows():
        name = str(row.get(name_col, "")).strip()
        sid = str(row.get(id_col, "")).strip()
        if name and sid and sid.lower() != "nan":
            students.append({"name": name, "student_number": sid})

    if log_func:
        log_func(f"Loaded {len(students)} students successfully.")
    return students


def kill_word():
    if sys.platform.startswith("win"):
        os.system("taskkill /f /im WINWORD.EXE >nul 2>&1")
        time.sleep(1)


def convert_with_word(docx_path, pdf_path, log_func=None):
    """Convert DOCX to PDF using Microsoft Word (Windows only)"""
    if not sys.platform.startswith("win"):
        if log_func:
            log_func("⚠️ Microsoft Word conversion only available on Windows")
        return False

    try:
        import win32com.client
        import pythoncom

        if log_func:
            log_func(
                f"📤 Converting with Microsoft Word: {os.path.basename(docx_path)}"
            )

        # Initialize COM
        pythoncom.CoInitialize()

        # Create Word application
        word = None
        doc = None

        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False

            # Open document
            doc = word.Documents.Open(os.path.abspath(docx_path))

            # Save as PDF
            doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)  # 17 = PDF format

            # Close document
            doc.Close()

            success = True
            if log_func:
                log_func(
                    f"✅ Microsoft Word conversion successful: {os.path.basename(pdf_path)}"
                )

        except Exception as e:
            if log_func:
                log_func(f"❌ Microsoft Word conversion error: {str(e)[:100]}")
            success = False

        finally:
            # Clean up
            if doc:
                try:
                    doc.Close()
                except:
                    pass
            if word:
                try:
                    word.Quit()
                except:
                    pass
            pythoncom.CoUninitialize()

        return success

    except ImportError:
        if log_func:
            log_func("⚠️ win32com not available for Microsoft Word conversion")
        return False
    except Exception as e:
        if log_func:
            log_func(f"❌ Microsoft Word error: {str(e)[:100]}")
        return False


def create_simple_pdf(pdf_path, student_data, log_func=None):
    """Create a simple PDF as last resort"""
    try:
        # Try to import reportlab
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
            from reportlab.lib.units import inch

            has_reportlab = True
        except ImportError:
            has_reportlab = False

        if not has_reportlab:
            if log_func:
                log_func("⚠️ reportlab not available for simple PDF")
            return False

        if log_func:
            log_func(f"📤 Creating simple PDF: {os.path.basename(pdf_path)}")

        # Create PDF
        c = canvas.Canvas(pdf_path, pagesize=letter)
        width, height = letter

        # Set up fonts
        c.setFont("Helvetica-Bold", 24)
        c.drawString(
            1 * inch,
            height - 1 * inch,
            f"Certificate for {student_data.get('name', '')}",
        )

        c.setFont("Helvetica", 14)

        c.drawString(
            1 * inch, height - 1.8 * inch, f"Date: {student_data.get('date', '')}"
        )

        c.setFont("Helvetica", 12)
        c.drawString(
            1 * inch,
            height - 2.5 * inch,
            "This document certifies successful completion",
        )
        c.drawString(1 * inch, height - 2.8 * inch, "of the course requirements.")

        # Add border
        c.rect(0.5 * inch, 0.5 * inch, width - 1 * inch, height - 1 * inch)

        c.save()

        if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
            if log_func:
                log_func(f"✅ Simple PDF created: {os.path.basename(pdf_path)}")
            return True
        else:
            if log_func:
                log_func(f"❌ Simple PDF creation failed: empty file")
            return False

    except Exception as e:
        if log_func:
            log_func(f"❌ Simple PDF creation error: {str(e)[:100]}")
        return False


# ---------------- Worker Class ----------------
class Worker(threading.Thread):
    def __init__(
        self,
        students,
        template_path,
        ui_callback,
        gen_docx=False,
        gen_pdf=False,
        merge_docx=False,
        merge_pdf=False,
    ):
        super().__init__()
        self.students = students
        self.template_path = template_path
        self.ui_callback = ui_callback
        self.gen_docx = gen_docx
        self.gen_pdf = gen_pdf
        self.merge_docx = merge_docx
        self.merge_pdf = merge_pdf

    def run(self):
        play_wav("assets/sounds/Generating.wav")
        total_steps = len(self.students)
        if self.gen_docx:
            total_steps += 1
        if self.gen_pdf:
            total_steps += 1
        if self.merge_docx:
            total_steps += 1
        if self.merge_pdf:
            total_steps += 1

        current_step = 0
        docx_files, pdf_files = [], []

        # ---------------- Generate DOCX files ----------------
        for s in self.students:
            try:
                tpl = DocxTemplate(self.template_path)
                context = {
                    "name": s["name"],
                    "student_number": s["student_number"],
                }
                placeholders_in_template = tpl.get_undeclared_template_variables()
                required = {"name", "student_number"}

                missing_required = required - placeholders_in_template
                if missing_required:
                    self.ui_callback(
                        progress=current_step / total_steps,
                        message=f"❌ Template missing required placeholder(s): {', '.join(missing_required)}",
                        error=True,
                    )
                    return

                tpl.render(context)

                safe_name = sanitize_filename(s["name"])
                filename = f"{s['student_number']}_{safe_name}.docx"
                out_docx = os.path.join(DOCX_OUT, filename)
                os.makedirs(os.path.dirname(out_docx), exist_ok=True)
                tpl.save(out_docx)
                docx_files.append(out_docx)

                current_step += 1
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"✅ Generated DOCX: {filename}",
                )

            except Exception as e:
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"❌ Error for {s['name']}: {str(e)[:150]}",
                    error=True,
                )
                return  # Stop immediately

            except Exception as e:
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"❌ Error for {s['name']}: {str(e)[:150]}",
                    error=True,
                )
                return  # Stop immediately

        # ---------------- Generate PDF files ----------------
        if self.gen_pdf:
            self.ui_callback(
                progress=current_step / total_steps,
                message="🔄 Starting PDF conversion...",
            )
            os.makedirs(PDF_OUT, exist_ok=True)

            total_files = len(docx_files)
            processed_files = 0

            for f in docx_files:
                pdf_name = os.path.splitext(os.path.basename(f))[0] + ".pdf"
                pdf_path = os.path.join(PDF_OUT, pdf_name)
                success = False

                # Find student data for this file
                student_data = None
                for s in self.students:
                    safe_name = sanitize_filename(s["name"])
                    if f"{s['student_number']}_{safe_name}.docx" in f:
                        student_data = {
                            "name": s["name"],
                            "student_number": s["student_number"],
                        }
                        break

                # Method 1: Microsoft Word (Windows only)
                if sys.platform.startswith("win"):
                    self.ui_callback(message=f"🔄 Trying Microsoft Word: {pdf_name}")
                    success = convert_with_word(
                        f, pdf_path, lambda msg: self.ui_callback(message=msg)
                    )
                    if success and is_valid_pdf(pdf_path):
                        pdf_files.append(pdf_path)
                        self.ui_callback(
                            message=f"✅ PDF created (Microsoft Word): {pdf_name}"
                        )
                    else:
                        success = False

                # Method 2: Simple PDF fallback
                if not success and student_data:
                    self.ui_callback(message=f"🔄 Creating simple PDF for: {pdf_name}")
                    success = create_simple_pdf(
                        pdf_path,
                        student_data,
                        lambda msg: self.ui_callback(message=msg),
                    )
                    if success and os.path.exists(pdf_path):
                        pdf_files.append(pdf_path)
                        self.ui_callback(message=f"✅ Simple PDF created: {pdf_name}")
                    else:
                        success = False

                # If all PDF methods fail
                if not success:
                    self.ui_callback(
                        message=f"❌ PDF conversion failed for: {pdf_name}", error=True
                    )
                    return  # Stop immediately

                processed_files += 1
                progress_increment = 1 / total_files / total_steps
                current_progress = current_step / total_steps + (
                    processed_files * progress_increment
                )
                self.ui_callback(
                    progress=current_progress,
                    message=f"🔄 Converting ({processed_files}/{total_files})...",
                )
                time.sleep(0.3)  # small delay

            current_step += 1
            self.ui_callback(
                progress=current_step / total_steps,
                message="✅ PDF conversion completed",
            )
            kill_word()  # cleanup

        # ---------------- Merge DOCX files ----------------
        if self.merge_docx and docx_files:
            try:
                master = Document(docx_files[0])
                composer = Composer(master)
                for f in docx_files[1:]:
                    composer.append(Document(f))

                merged_path = os.path.join(MERGED_DOCX_OUT, "MERGED_ALL.docx")
                os.makedirs(MERGED_DOCX_OUT, exist_ok=True)
                composer.save(merged_path)

                current_step += 1
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"✅ Merged DOCX saved: {merged_path}",
                )
            except Exception as e:
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"❌ Error merging DOCX: {str(e)[:150]}",
                    error=True,
                )
                return

        # ---------------- Merge PDF files ----------------
        if self.merge_pdf:
            if not self.gen_pdf:
                self.ui_callback(
                    message="⚠️ Cannot merge PDFs without generating PDFs.", error=True
                )
                return
            try:
                valid_pdfs = [f for f in pdf_files if is_valid_pdf(f)]
                if not valid_pdfs:
                    self.ui_callback(message="⚠️ No valid PDFs to merge.", error=True)
                    return

                merger = PdfMerger(strict=False)
                for f in valid_pdfs:
                    merger.append(f)
                merged_pdf_path = os.path.join(MERGED_PDF_OUT, "MERGED_ALL.pdf")
                os.makedirs(MERGED_PDF_OUT, exist_ok=True)
                merger.write(merged_pdf_path)
                merger.close()
                current_step += 1
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"✅ Merged PDF saved: {merged_pdf_path}",
                )

            except Exception as e:
                self.ui_callback(
                    progress=current_step / total_steps,
                    message=f"❌ Error merging PDFs: {str(e)[:150]}",
                    error=True,
                )
                return

        self.ui_callback(progress=1.0, message="✅ All tasks completed.", done=True)
        play_wav("assets/sounds/completed.wav")


class FAMSApp:
    def __init__(self, root):
        global topf
        self.root, self.students, self.template_path = root, [], None
        self.all_buttons = []
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        root.title("FAMS - Form Automation Management System")
        root.geometry("950x680")
        root.iconbitmap("assets/mbc.ico")
        root.configure(bg="#f0f4ff")

        style = ttk.Style(root)
        style.theme_use("clam")

        style.configure(
            "Amazing.Treeview",
            background="#ffffff",
            foreground="#333333",
            rowheight=28,
            fieldbackground="#ffffff",
            font=("Segoe UI", 10),
        )

        style.configure(
            "Amazing.Treeview.Heading",
            background="#2e4a9d",
            foreground="white",
            font=("Segoe UI", 11, "bold"),
            relief="raised",
        )

        style.map(
            "Amazing.Treeview",
            background=[("selected", "#c7d7ff")],
            foreground=[("selected", "#000000")],
        )

        style.map("Amazing.Treeview.Heading", background=[("active", "#1f3a8a")])

        style.configure(
            "Red.Horizontal.TProgressbar",
            troughcolor="#f0f0f0",
            background="#dd0f08",  # RED
            thickness=18,
        )

        style.configure(
            "TButton",
            font=("Segoe UI", 11, "bold"),
            foreground="white",
            background="#2e8b57",
            padding=6,
        )

        header_frame = tk.Frame(root, bg="#2e4a9d")
        header_frame.pack(fill="x")
        tk.Label(
            header_frame,
            text="Form Automation Management System ",
            font=("Segoe UI", 18, "bold"),
            bg="#2e4a9d",
            fg="white",
            pady=10,
        ).pack(side="left", padx=15)
        tk.Label(
            header_frame,
            text="Metro Business College",
            font=("Segoe UI", 10),
            bg="#2e4a9d",
            fg="white",
        ).place(x=20, y=60)
        try:
            img = Image.open("assets/mbc.png").resize((90, 90))
            self.logo_img = ImageTk.PhotoImage(img)
            tk.Label(header_frame, image=self.logo_img, bg="#2e4a9d").pack(
                side="right", padx=15
            )
        except Exception:
            tk.Label(
                header_frame,
                text="[Logo]",
                font=("Segoe UI", 12, "italic"),
                fg="white",
                bg="#2e4a9d",
            ).pack(side="right", padx=15)
        topf = tk.Frame(root, bg="#f0f4ff")
        topf.pack(fill="x", padx=12, pady=10)
        topf.grid_columnconfigure(3, weight=1)
        tk.Label(
            topf,
            text="① Upload Student File (CSV/Excel):",
            font=("Segoe UI", 11, "bold"),
            bg="#f0f4ff",
            fg="#2e4a9d",
        ).grid(row=0, column=0, sticky="w")
        self.lbl_file = tk.Label(
            topf, text="No file selected", fg="red", bg="#f0f4ff", font=("Segoe UI", 10)
        )
        self.lbl_file.grid(row=0, column=1, sticky="w", padx=8)
        browse_img = Image.open("assets/browse.png").resize((20, 20))
        self.browse_imgtk = ImageTk.PhotoImage(browse_img)
        self.btn_file = ctk.CTkButton(
            topf,
            text="Browse",
            font=("Segoe UI", 13, "bold"),
            image=self.browse_imgtk,
            corner_radius=5,
            fg_color=green_hex,
            hover_color=hover_green_hex,
            cursor="hand2",
            width=110,
            height=30,
            command=self.browse_file,
        )
        self.btn_file.grid(row=0, column=2, padx=6)
        self.all_buttons.append(self.btn_file)

        tk.Label(
            topf,
            text="② Select DOCX Template:",
            font=("Segoe UI", 11, "bold"),
            bg="#f0f4ff",
            fg="#2e4a9d",
        ).grid(row=1, column=0, sticky="w", pady=6)
        self.lbl_template = tk.Label(
            topf,
            text="No template selected",
            fg="red",
            bg="#f0f4ff",
            font=("Segoe UI", 10),
        )
        self.lbl_template.grid(row=1, column=1, sticky="w")
        self.btn_template = ctk.CTkButton(
            topf,
            text="Browse",
            font=("Segoe UI", 13, "bold"),
            image=self.browse_imgtk,
            corner_radius=5,
            fg_color=green_hex,
            hover_color=hover_green_hex,
            cursor="hand2",
            width=110,
            height=30,
            command=self.browse_template,
        )
        self.btn_template.grid(row=1, column=2, padx=6)
        self.all_buttons.append(self.btn_template)

        help_img = Image.open("assets/help.png").resize((20, 20))
        help_imgtk = ImageTk.PhotoImage(help_img)
        ctk.CTkButton(
            topf,
            text="Help",
            image=help_imgtk,
            cursor="hand2",
            font=("Segoe UI", 15, "bold"),
            text_color="#FFFFFF",
            width=10,
            command=self.show_help,
        ).grid(row=0, column=3, sticky="e", padx=10)

        midf = tk.Frame(root, bg="#e1e5f0")
        midf.pack(fill="both", expand=True, padx=12, pady=6)
        left = tk.LabelFrame(
            midf,
            text="Loaded Students",
            bg="#f0f4ff",
            fg="#2e4a9d",
            font=("Segoe UI", 11, "bold"),
        )
        left.pack(side="left", fill="both", expand=True, padx=(0, 6))
        scrollbar = ttk.Scrollbar(left, orient="vertical")
        scrollbar.pack(side="right", fill="y")
        self.tree = ttk.Treeview(
            left,
            columns=("sid", "name"),
            show="headings",
            height=14,
            style="Amazing.Treeview",
            yscrollcommand=scrollbar.set,
        )
        self.tree.tag_configure("oddrow", background="#f4f6fb")
        self.tree.tag_configure("evenrow", background="#ffffff")
        self.tree.heading("sid", text="Student Number")
        self.tree.heading("name", text="Student Name")
        self.tree.column("sid", width=160, anchor="center")
        self.tree.column("name", width=360, anchor="w")
        self.tree.pack(fill="both", expand=True, padx=6, pady=4)

        scrollbar.config(command=self.tree.yview)
        self.lbl_count = tk.Label(
            left,
            text="0 students loaded",
            bg="#f0f4ff",
            fg="black",
            font=("Segoe UI", 10, "italic"),
        )
        self.lbl_count.pack(anchor="w", pady=4)
        right = tk.LabelFrame(
            midf,
            text="Actions & Log",
            bg="#f0f4ff",
            fg="#2e4a9d",
            font=("Segoe UI", 11, "bold"),
        )
        right.pack(side="right", fill="y")

        gen_img = Image.open("assets/genrate.png").resize((30, 30))
        gen_imgtk = ImageTk.PhotoImage(gen_img)
        self.generate_btn = ctk.CTkButton(
            right,
            text="Generate Documents",
            image=gen_imgtk,
            cursor="hand2",
            font=("Segoe UI", 15, "bold"),
            text_color="#FFFFFF",
            width=10,
            command=self.start_generate,
        )
        self.generate_btn.pack(pady=(8, 6))
        self.all_buttons.append(self.generate_btn)

        self.docx_var = tk.BooleanVar()
        self.pdf_var = tk.BooleanVar()
        self.merge_docx_var = tk.BooleanVar()
        self.merge_pdf_var = tk.BooleanVar()
        ctk.CTkCheckBox(
            right,
            text="Generate Docx",
            variable=self.docx_var,
            text_color="#000000",
            font=("Segoe UI", 14, "bold"),
        ).pack(pady=2, padx=50, anchor="w")

        ctk.CTkCheckBox(
            right,
            text="Generate PDF also",
            variable=self.pdf_var,
            text_color="#000000",
            font=("Segoe UI", 14, "bold"),
        ).pack(pady=2, padx=50, anchor="w")

        ctk.CTkCheckBox(
            right,
            text="Merge all DOCX",
            variable=self.merge_docx_var,
            text_color="#000000",
            font=("Segoe UI", 14, "bold"),
        ).pack(pady=2, padx=50, anchor="w")

        ctk.CTkCheckBox(
            right,
            text="Merge all PDFs",
            variable=self.merge_pdf_var,
            text_color="#000000",
            font=("Segoe UI", 14, "bold"),
        ).pack(pady=2, padx=50, anchor="w")

        openfolder_img = Image.open("assets/openfolder.png").resize((30, 30))
        openfoldertk = ImageTk.PhotoImage(openfolder_img)

        self.btn_open_output = ctk.CTkButton(
            right,
            text="Open Output Folder",
            image=openfoldertk,
            command=self.open_output,
            font=("Segoe UI", 12, "bold"),
            text_color="#FFFFFF",
            cursor="hand2",
        )
        self.btn_open_output.pack(pady=6)
        self.all_buttons.append(self.btn_open_output)

        btn_frame = tk.Frame(right, bg="#f0f4ff")
        btn_frame.pack(pady=6, fill="x")

        img_log = Image.open("assets/log.png").resize((20, 20))
        img_logtk = ImageTk.PhotoImage(img_log)
        self.btn_logs = ctk.CTkButton(
            btn_frame,
            text="Download Logs",
            font=("Segoe UI", 12, "bold"),
            fg_color="#7C0E0E",
            command=self.save_logs,
            cursor="hand2",
            image=img_logtk,
        )
        self.btn_logs.pack(side="left", expand=True, fill="x", padx=(3, 3))
        self.all_buttons.append(self.btn_logs)

        clear_img = Image.open("assets/clear.png").resize((20, 20))
        clear_imgtk = ImageTk.PhotoImage(clear_img)
        self.btn_clear = ctk.CTkButton(
            btn_frame,
            text="Clear Fields",
            image=clear_imgtk,
            command=self.clear_fields,
            fg_color="#EE0808",
            font=("Segoe UI", 12, "bold"),
            width=110,
        )
        self.btn_clear.pack(
            side="left",
            expand=True,
            fill="x",
            padx=(5, 5),
        )
        self.all_buttons.append(self.btn_clear)

        self.progress = ttk.Progressbar(
            right, length=260, mode="determinate", style="Red.Horizontal.TProgressbar"
        )
        self.progress.pack(pady=8)
        self.progress_label = tk.Label(
            right, text="0%", bg="#f0f4ff", font=("Segoe UI", 10, "bold")
        )
        self.progress_label.pack()
        tk.Label(
            right,
            text="Activity Log:",
            bg="#f0f4ff",
            fg="black",
            font=("Segoe UI", 10, "bold"),
        ).pack()
        self.log = tk.Text(
            right, width=36, height=12, state="disabled", font=("Consolas", 9)
        )
        self.log.pack(padx=6, pady=(0, 6))

    def on_close(self):
        if messagebox.askokcancel("Exit", "Are you sure to exit FAMS?"):
            try:
                os._exit(0)  # HARD EXIT – guaranteed
            except:
                self.root.destroy()

    def browse_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xls;*.xlsx")]
        )
        if not path:
            return
        try:
            students = read_students(path)
            self.students = students
            self.lbl_file.configure(text=os.path.basename(path), fg="green")
            self.btn_file.configure(
                text="Browse",
            )

            self.check_browse = Image.open("assets/check.png").resize((30, 30))
            self.check_browse_tk = ImageTk.PhotoImage(self.check_browse)
            self.check_lbl_browse = tk.Label(
                topf, image=self.check_browse_tk, bg="#f0f4ff"
            )
            self.check_lbl_browse.grid(row=0, column=3, sticky="w")
            self.refresh_table()
            self.log_message(f"✅ Loaded {len(students)} students")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read file: {e}")
            print(f"Failed to read file: {e}")

    def browse_template(self):
        path = filedialog.askopenfilename(filetypes=[("Word Document", "*.docx")])
        if path and path.lower().endswith(".docx"):
            # Load template
            try:
                tpl = DocxTemplate(path)
                required_placeholders = {"name", "student_number"}
                template_placeholders = tpl.get_undeclared_template_variables()
                missing = required_placeholders - template_placeholders
                if missing:
                    messagebox.showerror(
                        "Template Error",
                        f"The template is missing required placeholders:\n{', '.join(missing)}",
                    )
                    return
            except Exception as e:
                messagebox.showerror("Template Error", f"Failed to read template: {e}")
                return

            self.template_path = path
            self.lbl_template.configure(text=os.path.basename(path), fg="green")
            self.btn_template.configure(text="Browse")
            self.check_temp = Image.open("assets/check.png").resize((30, 30))
            self.check_temp_tk = ImageTk.PhotoImage(self.check_temp)
            self.check_lbl_temp = tk.Label(topf, image=self.check_temp_tk, bg="#f0f4ff")
            self.check_lbl_temp.grid(row=1, column=3, sticky="w")
            self.log_message(f"✅ Template selected: {os.path.basename(path)}")
        else:
            messagebox.showerror("Invalid", "Please select a DOCX template.")

    def speak(self, text):
        engine = pyttsx3.init()
        voices = engine.getProperty("voices")

        for voice in voices:
            if "female" in voice.name.lower() or "zira" in voice.name.lower():
                engine.setProperty("voice", voice.id)
                break

        engine.say(text)
        engine.runAndWait()

    def refresh_table(self):
        for r in self.tree.get_children():
            self.tree.delete(r)

        for i, s in enumerate(self.students[:2000]):
            tag = "evenrow" if i % 2 == 0 else "oddrow"
            self.tree.insert(
                "", "end", values=(s["student_number"], s["name"]), tags=(tag,)
            )
        threading.Thread(
            target=lambda: self.speak(f"{len(self.students)} students loaded"),
            daemon=True,
        ).start()
        self.lbl_count.config(text=f"{len(self.students)} students loaded")

    def start_generate(self):

        if not self.students:
            threading.Thread(
                target=lambda: play_wav("assets/sounds/upload_data_first.wav"),
                daemon=True,
            ).start()
            messagebox.showwarning("No data", "Upload student data first.")
            return
        if not self.template_path:
            threading.Thread(
                target=lambda: play_wav("assets/sounds/select_docx_temp.wav"),
                daemon=True,
            ).start()
            messagebox.showwarning("No template", "Select a DOCX template.")
            return
        if not (
            self.docx_var.get()
            or self.pdf_var.get()
            or self.merge_docx_var.get()
            or self.merge_pdf_var.get()
        ):
            messagebox.showwarning(
                "No options selected",
                "Please check at least one option to generate:\n"
                "- Generate Docx\n"
                "- Generate PDF also\n"
                "- Merge all DOCX\n"
                "- Merge all PDFs",
            )
            return

        for folder in [DOCX_OUT, PDF_OUT, MERGED_DOCX_OUT, MERGED_PDF_OUT]:
            os.makedirs(folder, exist_ok=True)
        self.progress["value"] = 0
        self.progress_label.config(text="0%")
        self.log_message("🔄 Starting document generation...")

        Worker(
            self.students,
            self.template_path,
            self.worker_callback,
            self.docx_var.get(),
            self.pdf_var.get(),
            self.merge_docx_var.get(),
            self.merge_pdf_var.get(),
        ).start()

    def worker_callback(self, progress=0.0, message="", done=False, error=False):
        """
        Callback function from Worker thread.
        - progress: float (0.0 to 1.0)
        - message: string log message
        - done: True if all tasks completed
        - error: True if an error occurred
        """

        def update():
            # Update progress bar
            self.progress["value"] = progress * 100
            self.progress_label.config(text=f"{int(progress*100)}%")

            # Log message
            if message:
                self.log_message(message)

            # Show error popup if error flag is set
            if error and message:
                messagebox.showerror(
                    "Error Occurred", f"{message}\n\nSee Activity Log for details."
                )

            # Show completion info
            if done:

                messagebox.showinfo("Done", "All tasks completed successfully.")

        self.root.after(0, update)

    def log_message(self, text):
        ts = datetime.now().strftime("%H:%M:%S")
        msg = f"[{ts}] {text}\n"
        self.log.config(state="normal")
        self.log.insert("end", msg)
        self.log.see("end")
        self.log.config(state="disabled")

    def save_logs(self):
        os.makedirs(BASE_OUT, exist_ok=True)
        logs = self.log.get("1.0", "end").strip()
        if not logs:
            messagebox.showinfo("Logs", "No logs to save.")
            return
        with open(LOG_FILE, "w", encoding="utf-8") as f:
            f.write(logs)
        threading.Thread(
            target=lambda: play_wav("assets/sounds/download_log.wav"), daemon=True
        ).start()
        messagebox.showinfo("Logs Saved", f"Logs saved to:\n{LOG_FILE}")

    def show_help(self):
        help_win = tk.Toplevel(self.root)
        help_win.title("FAMS – Help Guide")
        help_win.geometry("600x600")
        help_win.resizable(False, False)
        help_win.transient(self.root)
        help_win.iconbitmap("assets/mbc.ico")
        help_win.grab_set()
        logo_path = "assets/mbc.png"

        frame = tk.Frame(help_win, padx=15, pady=15)
        frame.pack(fill="both", expand=True)
        try:
            logo_img = Image.open(logo_path)
            logo_img = logo_img.resize((100, 100), Image.LANCZOS)
            self.help_logo = ImageTk.PhotoImage(logo_img)

            logo_label = tk.Label(help_win, image=self.help_logo, borderwidth=0)
            logo_label.place(relx=1, x=-30, y=30, anchor="ne")
        except Exception as e:
            print(f"Failed to load logo: {e}")

        canvas = tk.Canvas(frame, borderwidth=0)
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=canvas.yview)
        scroll_frame = tk.Frame(canvas)

        scroll_frame.bind(
            "<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        def section_title(text):
            return tk.Label(
                scroll_frame,
                text=text,
                font=("Segoe UI", 14, "bold"),
                fg="#2e4a9d",
                anchor="w",
                pady=8,
            )

        def bullet_point(text, icon=None, fg="black"):
            frame = tk.Frame(scroll_frame, pady=2)
            frame.pack(anchor="w", fill="x")
            if icon:
                lbl_icon = tk.Label(
                    frame, text=icon, font=("Segoe UI Emoji", 12), fg=fg
                )
                lbl_icon.pack(side="left")
            lbl_text = tk.Label(
                frame, text=text, font=("Segoe UI", 11), justify="left", wraplength=480
            )
            lbl_text.pack(side="left", padx=6)
            return frame

        def ssexample(event=None):
            image_path = "assets/ss_example.png"
            if not os.path.exists(image_path):
                return

            try:
                img = Image.open(image_path)
                img.show()
            except:
                try:
                    if sys.platform == "win32":
                        os.startfile(image_path)
                    elif sys.platform == "darwin":
                        os.system(f"open {image_path}")
                    else:
                        os.system(f"xdg-open {image_path}")
                except:
                    pass

        def ss1example(event=None):
            image_path = "assets/ss1_example.png"
            if not os.path.exists(image_path):
                return

            try:
                img = Image.open(image_path)
                img.show()
            except:
                try:
                    if sys.platform == "win32":
                        os.startfile(image_path)
                    elif sys.platform == "darwin":
                        os.system(f"open {image_path}")
                    else:
                        os.system(f"xdg-open {image_path}")
                except:
                    pass

        section_title("1. Upload Student File").pack(fill="x")
        bullet_point("Click 'Browse' under Upload Student File", icon="☑️").pack()
        bullet_point("Supported formats: CSV, XLS, XLSX").pack()
        bullet_point(
            "File must contain Name and Student Number columns\n(𝐞.𝐠 𝐬𝐭𝐮𝐝𝐞𝐧𝐭_𝐧𝐮𝐦𝐛𝐞𝐫 | 𝐧𝐚𝐦𝐞)"
        ).pack()
        bullet_point("Format Supported:").pack()
        bullet_point(
            '𝐅𝐨𝐫 𝐒𝐭𝐮𝐝𝐞𝐧𝐭 𝐍𝐮𝐦𝐛𝐞𝐫: \n["student_number","student no","student_no","id","studentid",\n"student number"]'
        ).pack()
        bullet_point(
            '𝐅𝐨𝐫 𝐒𝐭𝐮𝐝𝐞𝐧𝐭 𝐍𝐚𝐦𝐞: ["name", "full name", "student name", "student"]'
        ).pack()
        example = tk.Label(
            scroll_frame,
            text="Screen Shot Example",
            font=("Segoe UI", 11, "underline"),
            fg="blue",
            cursor="hand2",
        )
        example.pack(anchor="w", pady=(0, 6))
        example.bind("<Button-1>", ssexample)

        section_title("2. Select DOCX Template").pack(fill="x")
        bullet_point("Choose a Word (.docx) template", icon="☑️").pack()
        bullet_point(
            "⚠️ Make sure use placeholders like {{ 𝐧𝐚𝐦𝐞 }}, {{ 𝐬𝐭𝐮𝐝𝐞𝐧𝐭_𝐧𝐮𝐦𝐛𝐞𝐫 }}"
        ).pack()
        example1 = tk.Label(
            scroll_frame,
            text="Screen Shot Example",
            font=("Segoe UI", 11, "underline"),
            fg="blue",
            cursor="hand2",
        )
        example1.pack(anchor="w", pady=(0, 6))
        example1.bind("<Button-1>", ss1example)

        section_title("📄 Template Page Break & Formatting").pack(fill="x")
        bullet_point(
            "Each student record generates ONE document or ONE page", icon="ℹ️"
        ).pack()
        bullet_point(
            "Use a manual Page Break (Ctrl + Enter) to separate pages", icon="📌"
        ).pack()
        bullet_point(
            "Do NOT add extra blank pages at the end of the template",
            icon="⚠️",
            fg="#a94442",
        ).pack()
        bullet_point(
            "Place all placeholders on the same page\n(e.g {{ name }}, {{ student_number }})"
        ).pack()
        bullet_point(
            "Avoid putting placeholders inside text boxes, shapes, or headers unless required",
            icon="🚫",
        ).pack()
        bullet_point(
            "If merging DOCX or PDF files, page breaks control final document layout",
            icon="🧩",
        ).pack()
        bullet_point(
            "Headers and footers are supported, but must contain valid placeholders only",
            icon="✔️",
        ).pack()

        section_title("3. Generate Documents").pack(fill="x")
        bullet_point("Click 🚀 Generate Documents", icon="▶️").pack()
        bullet_point("Optional:", fg="#555555").pack()
        bullet_point("Generate Docx", icon="✓").pack()
        bullet_point("Generate PDF also", icon="✓").pack()
        bullet_point("Merge all DOCX", icon="✓").pack()
        bullet_point("Merge all PDFs", icon="✓").pack()

        section_title("4. PDF Generation (Priority Order)").pack(fill="x")
        bullet_point("1. Microsoft Word (Windows only, best quality)", icon="①").pack()
        # bullet_point("2. LibreOffice Portable (if Word fails)", icon="②").pack()
        bullet_point("2. Simple PDF (last resort)", icon="②").pack()
        # bullet_point(
        #     "✅ Uses local LibreOffice from 'LibreOffice/program/soffice.exe' if needed",
        #     icon="💻",
        # ).pack()
        # bullet_point(
        #     "Ensure LibreOffice folder exists in the same directory as FAMS", icon="📁"
        # ).pack()
        # bullet_point(
        #     "Close all Word/LibreOffice windows before running PDF conversion", icon="⚠️"
        # ).pack()

        section_title("5. Output Files").pack(fill="x")
        bullet_point(
            "Generated files are saved in the 'fams_output' folder", icon="📁"
        ).pack()

        def open_output_folder(event=None):
            self.open_output()

        output_link = tk.Label(
            scroll_frame,
            text="Open Output Folder",
            font=("Segoe UI", 11, "underline"),
            fg="blue",
            cursor="hand2",
        )
        output_link.pack(anchor="w", pady=(0, 6))
        output_link.bind("<Button-1>", open_output_folder)

        section_title("6. Logs").pack(fill="x")
        bullet_point("All actions appear in the Activity Log").pack()

        def save_logs(event=None):
            self.save_logs()

        logs_link = tk.Label(
            scroll_frame,
            text="Download Logs",
            font=("Segoe UI", 11, "underline"),
            fg="blue",
            cursor="hand2",
        )
        logs_link.pack(anchor="w", pady=(0, 10))
        logs_link.bind("<Button-1>", save_logs)

        section_title("🛠 Tips").pack(fill="x")
        bullet_point("Do not open DOCX files while generating PDFs").pack()
        bullet_point("Make sure the template file is closed").pack()
        bullet_point("Large student lists may take time").pack()
        # bullet_point(
        #     "Close all Word/LibreOffice windows before running", icon="⚠️"
        # ).pack()
        # bullet_point(
        #     "For best PDF quality, ensure LibreOffice folder exists", icon="✅"
        # ).pack()
        bullet_point(
            "Windows users get best results with Microsoft Word", icon="💯"
        ).pack()

        final_lbl = tk.Label(
            scroll_frame,
            text="✅ You're good to go!",
            font=("Segoe UI", 12, "bold"),
            fg="#2e8b57",
            pady=10,
        )
        final_lbl.pack(anchor="w")

        ttk.Button(
            help_win, text="Close", command=help_win.destroy, cursor="hand2"
        ).pack(pady=10)

    def open_output(self):
        path = os.path.abspath(BASE_OUT)
        os.makedirs(path, exist_ok=True)
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])

    def clear_fields(self):
        self.students = []
        self.template_path = None
        self.lbl_file.config(text="No file selected", fg="red")
        self.lbl_template.config(text="No template selected", fg="red")

        if hasattr(self, "check_lbl_temp") and self.check_lbl_temp.winfo_ismapped():
            self.check_lbl_temp.grid_forget()

        if hasattr(self, "check_lbl_browse") and self.check_lbl_browse.winfo_ismapped():
            self.check_lbl_browse.grid_forget()

        self.btn_template.configure(text="Browse")

        for r in self.tree.get_children():
            self.tree.delete(r)
        self.lbl_count.config(text="0 students loaded")
        self.log.config(state="normal")
        self.log.delete("1.0", "end")
        self.log.config(state="disabled")
        self.progress["value"] = 0
        self.progress_label.config(text="0%")
        threading.Thread(
            target=lambda: play_wav("assets/sounds/Cleared.wav"), daemon=True
        ).start()
        self.log_message("🧹 Cleared all fields.")


if __name__ == "__main__":
    splash = tk.Tk()
    splash.overrideredirect(True)
    splash.wm_attributes("-topmost", True)
    threading.Thread(
        target=lambda: play_wav("assets/sounds/fams.wav"), daemon=True
    ).start()

    if sys.platform.startswith("win"):
        splash.configure(bg="#FF66C4")
        splash.wm_attributes("-transparentcolor", "#FF66C4")

    try:
        logo_img = Image.open("assets/splash.png").resize((600, 600))
        splash.logo = ImageTk.PhotoImage(logo_img)
        width, height = splash.logo.width(), splash.logo.height()
    except:
        splash.logo = None
        width, height = 200, 200

    screen_w = splash.winfo_screenwidth()
    screen_h = splash.winfo_screenheight()
    x = (screen_w // 2) - (width // 2)
    y = (screen_h // 2) - (height // 2)
    splash.geometry(f"{width}x{height}+{x}+{y}")

    label = tk.Label(
        splash, image=splash.logo, bg="#FF66C4", borderwidth=0, highlightthickness=0
    )
    label.pack()

    splash.update()
    splash.after(3000, splash.destroy)
    splash.mainloop()

    root = tk.Tk()
    app = FAMSApp(root)
    root.mainloop()
